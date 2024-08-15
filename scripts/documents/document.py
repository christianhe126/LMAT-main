import json
import base64
import sys
from io import BytesIO
import re
from bs4 import BeautifulSoup
import textwrap
import math
from PIL import Image
from pptx.util import Inches, Pt
import documents.bullet as bullet

class Document:
    def __init__(self, format, data, template=None, metadata=None):
        self.format = format
        self.html = data['html']
        self.title = "Presentation"
        self.metadata = metadata

        if 'title' in data:
            self.title = data['title']
        if self.metadata and 'Title' in self.metadata:
            self.title = self.metadata['Title']

        if template:
            self.path = "./resources/customTemplate" + format
            self._write_base64_to_file(template, self.path)
        else:
            self.path = "./resources/template" + format

    def process(self):
        self.html.replace('\n', '')

        soup = BeautifulSoup(self.html, 'html.parser')
        first_tag = soup.find()
        if first_tag and first_tag.name != 'h1':
            new_h1 = soup.new_tag('h1')
            new_h1.string = "[No Title]"
            soup.insert(0, new_h1)
        
        paragraphs = []
        title = None
        for element in soup.contents:
            if hasattr(element, 'name'):
                if element.name == 'h1':
                    if title:
                        self.write(title, paragraphs)
                        paragraphs = []
                    title = element.get_text()

                elif element.name == 'table':
                    paragraph = self.Paragraph(str(element), type='table')
                    paragraphs.append(paragraph)

                elif element.name == 'img' or (element.name == 'figure' and element.find('img')) or (element.name == 'p' and element.find('img')):
                    img = None
                    if hasattr(element.find('img'), 'src'):
                        img = element.find('img')
                    elif hasattr(element, 'src'):
                        img = element

                    if img:
                        src = img['src']
                        estimate = None
                        if hasattr(img, 'width') and hasattr(img, 'height'):
                            estimate = min(int(hasattr(img, 'height') / 25), 1)

                        b64_string = src
                        if b64_string.startswith("data:image/"):
                            b64_string = b64_string.split(",")[1]
                        
                        image_data = base64.b64decode(b64_string)
                        ratio = None
                        with Image.open(BytesIO(image_data)) as img:
                            img_width, img_height = img.size
                            ratio = img_width / img_height
                            if ratio < 1:
                                estimate = 10
                        
                        image = self.Paragraph(src, type='img', estimate=estimate)
                        paragraphs.append(image)

                elif element.name == 'ul' or element.name == 'ol':
                    estimate = 0
                    if hasattr(self, "chars_per_line") and self.chars_per_line:
                        for li in element.contents:
                            if hasattr(li, 'name') and li.name and li.name == 'li':
                                wrap = textwrap.wrap(li.get_text(), width=self.chars_per_line)
                                estimate += len(wrap)
                    else:
                        estimate = len(element.find_all('li'))

                    paragraph = self.Paragraph(str(element), type='p', estimate=estimate)
                    paragraphs.append(paragraph)
                                      
                else:
                    paragraph = self.Paragraph(str(element), type='p')
                    paragraphs.append(paragraph)

        if title:
            self.write(title, paragraphs)
            title = None

        document_io = BytesIO()
        self.document.save(document_io)
        document_io.seek(0)
        return base64.b64encode(document_io.read()).decode('utf-8')
        

    def create_page(self, title = None):
        raise NotImplementedError("Subclasses must implement this method")

    def add_text(self, text_frame, html_text):
        raise NotImplementedError("Subclasses must implement this method")
    
    def add_image(self, page, x, y, width, height, source):
        raise NotImplementedError("Subclasses must implement this method")
    
    def write(self, title, paragraphs):
        # Reduce number of text elements
        self._reduce_paragraphs(paragraphs)

        # Process the paragraphs
        current_line_number = 0
        line_height = self.page_size["height"] / self.max_lines
        page = self.create_page(title=title)
        for paragraph in paragraphs:
            if paragraph.estimate >= self.max_lines or current_line_number + paragraph.estimate >= self.max_lines:
                page = self.create_page(title=title, following=True)
                current_line_number = 0
                
            x = int(self.page_size["x"])
            y = int(self.page_size["y"] + current_line_number*line_height)
            width = int(self.page_size["width"])
            height = int(line_height * paragraph.estimate)
            
            if paragraph.type == "p":
                self.add_text(page,x,y,width,height,paragraph.content)
                
            elif paragraph.type == "img":
                self.add_image(page,x,y,width,height,paragraph.content)
                
            elif paragraph.type == "table":
                self.add_table(page,x,y,width,height,paragraph.content)

            current_line_number += paragraph.estimate

    def html_to_element(self, object, html, p = None, font = None):
        if re.compile(r'<[^>]+>').search(html):
            soup = BeautifulSoup(html, 'html.parser')
            for element in soup.contents:
                if hasattr(element, 'name') and element.name:
                    if element.name == 'br' or element.name == 'p':
                        p = object.add_paragraph()
                        self._format(element.decode_contents(), p)                      
                            
                    elif element.name == 'ul':
                        for li in element.contents:
                            if hasattr(li, 'name') and li.name and li.name == 'li':
                                p = object.add_paragraph()
                                self._add_bullet_up_to_second_level(object, li, p)
                    
                    elif element.name == 'ol':
                        for i, li in enumerate(element.contents):
                            if hasattr(li, 'name') and li.name and li.name == 'li':
                                p = object.add_paragraph()
                                if not hasattr(self, "_format_bullet_point"):
                                    bullet.pNumber(p, "Arial")
                                self._format(li.decode_contents(), p)

                    elif element.name == 'li':
                        p = object.add_paragraph()
                        self._add_bullet_up_to_second_level(object, element, p)

                    else:
                        p = object.add_paragraph()
                        self._format(str(element), p)    

                    if font:
                        p.font.size = Pt(font)
                    
                else:
                    run = p.add_run()
                    run.text = element.get_text()
        else:
            p = object.add_paragraph()
            run = p.add_run()
            run.text = html
    
    def _format(self, html, p):
        if re.compile(r'<[^>]+>').search(html):
            soup = BeautifulSoup(html, 'html.parser')
            for element in soup.contents:
                run = p.add_run()
                run.text = element.get_text()
                if hasattr(element, 'name') and element.name:
                    if element.name == 'strong':
                        run.font.bold = True
                    if element.name == 'em':
                        run.font.italic = True
                    if element.name == 'u':
                        run.font.underline = True
                    if element.name == 'a':
                        if hasattr(run, 'hyperlink'):
                            hlink = run.hyperlink
                            if hasattr(element, 'href'):
                                hlink.address = element['href']
                            else:
                                hlink.address = ''
                        elif hasattr(self, '_add_hyperlink_into_run'):
                            self._add_hyperlink_into_run(p, run, element['href'])
        else:
            run = p.add_run()
            run.text = html
        
    def _add_bullet_up_to_second_level(self, object, element, p):
        if not hasattr(self, "_format_bullet_point"):
            bullet.pBullet(p, "Arial")
            
        texts = []
        if element.find('li'):
            _temp = []
            for li in element.find_all('li'):
                _temp.append({"text": li.decode_contents(), "level": 2})
                li.decompose()

            texts.append({"text": element.decode_contents(), "level": 1})
            for _t in _temp:
                texts.append(_t)
            
        else:
            texts.append({"text": element.decode_contents(), "level": 1})

        for t in texts:
            _t = t["text"]
            if hasattr(self, "_format_bullet_point"):
                _t = self._format_bullet_point(t["text"])

            if t["level"] == 2:
                p = object.add_paragraph()
            p.level = t["level"]
            self._format(_t, p)
    
    def _calculate_img_size_and_position(self, x, y, width, height, image_data):
        with Image.open(BytesIO(image_data)) as img:
            img_width, img_height = img.size
        
        width_ratio = width / img_width
        height_ratio = height / img_height
        
        if width_ratio < height_ratio:
            new_width = width
            new_height = int(img_height * width_ratio)
            y += (height - new_height) / 2
        else:
            new_width = int(img_width * height_ratio)
            new_height = height
            x += (width - new_width) / 2
        
        return x, y, new_width, new_height

    def _get_table_data(self, html_table):
        soup = BeautifulSoup(html_table, 'html.parser')
        table = soup.find('table')
        if (table is None):
            return 0, 0, "normal", 0, None, [], []
        b_style = "normal"
        b_width = 1
        b_color = None
        cellspacing = None
        #if hasattr(table, 'style') and len(table.get('style').split('border-style: ')) > 1 and len(table.get('style').split('border-style: ')[1].split(';')) > 0:
        #    b_style = table.get('style').split('border-style: ')[1].split(';')[0]
        if hasattr(table, 'border') and table.get('border') is not None:
            b_width = table.get('border')
        
        if hasattr(table, 'style') and len(table.get('style').split('background-color: ')) > 1:
            styles = table.get('style').split('background-color: ')[1]
            b_color = styles.split(';')[0]

        if hasattr(table, 'cellspacing') and table.get('cellspacing') is not None:
            cellspacing = table.get('cellspacing')

        # Initialize the data array and count rows and columns
        data = []
        sizes = []
        rows = table.find_all('tr')
        n_row = len(rows)
        n_col = 0

        # Extract table data
        for row in rows:
            cells = row.find_all('td')
            if len(cells) > n_col:
                n_col = len(cells)
            row_data = []
            row_sizes = []  # To store size info for the current row
            for cell in cells:
                row_data.append(cell.decode_contents())
                width = cell.get('style', '').split('width: ')[-1].split('%;')[0] if 'width' in cell.get('style', '') else None
                if 'height' not in cell.get('style', ''):
                    height = None
                else:
                    height = cell.get('style', '').split('height: ')[-1].split(';')[0].strip()
                    height = float(height.replace('px', '')) if 'px' in height else float(height.replace('pt', '')) * 1.333 if 'pt' in height else None

                row_sizes.append((width, height))
            data.append(row_data)
            sizes.append(row_sizes)

        return n_row, n_col, b_style, b_width, b_color, cellspacing, data, sizes
    
    def _write_base64_to_file(self, b64_string, file_path):
        binary_data = base64.b64decode(b64_string)
        with open(file_path, 'wb') as file:
            file.write(binary_data)

    class Paragraph:
        def __init__(self, content, type="p", chars_per_line=90, estimate=None):
            self.content = content
            self.chars_per_line = chars_per_line
            self.type = type
            if self.type == "p":
                self.estimate = self._number_of_lines()
            elif self.type == "img":
                self.estimate = 7
            elif self.type == "table":
                self.estimate = 7

            if estimate and estimate > 0:
                self.estimate = estimate

        def _number_of_lines(self):
            content_non_html = self._get_non_html()
            if content_non_html:
                wrapped_text = textwrap.wrap(content_non_html, width=self.chars_per_line)
            else:
                wrapped_text = textwrap.wrap(self.content, width=self.chars_per_line)
            return len(wrapped_text)
        
        def _get_non_html(self):
            soup = BeautifulSoup(self.content, 'html.parser')
            return soup.get_text()
        



